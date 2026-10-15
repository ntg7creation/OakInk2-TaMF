from pathlib import Path
import json, hashlib
import numpy as np
import pandas as pd
from PIL import Image, ImageOps, ImageDraw
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'cet_findings_report'
RUN=ROOT.parent/'common/partial_batch/test_one__0399/cet/20261006T015545971089Z'
rows=json.loads((RUN/'motions.json').read_text())
summary=json.loads((RUN/'summary.json').read_text())
records=[]
for r in rows:
    assert r['stage']=='R' and r['status']=='success'
    z=np.load(RUN/r['arrays_path'],allow_pickle=False)
    p=z['dtw_frame_pairs']; gt=z['GT_curl_degrees']; pred=z['prediction_curl_degrees']
    e=np.abs(pred[p[:,0]]-gt[p[:,1]])
    assert np.allclose(e.mean(0),r['per_joint_degrees'])
    assert np.isclose(e.mean(),r['cet_degrees'])
    d={k:r[k] for k in ['original_motion_index','primitive','description','hand_side','cet_degrees','duration_seconds','W','sequence']}
    d['action']=r['primitive'].split(':')[0]
    for k,v in r['GT_curl_descriptors'].items(): d[k]=v['whole_hand']
    for i,f in enumerate(summary['finger_names']):
        d[f+'_cet']=r['per_finger_degrees'][i]
        d[f+'_depth']=r['GT_curl_descriptors']['mean_absolute_flexion_degrees']['per_finger'][i]
        d[f+'_relative_depth_pct']=100*d[f+'_cet']/d[f+'_depth']
    d['top_finger']=summary['finger_names'][int(np.argmax(r['per_finger_degrees']))]
    jj=np.argsort(r['per_joint_degrees'])[::-1][:3]
    d['top_joints']='; '.join(f"{summary['joint_names'][i]} {r['per_joint_degrees'][i]:.3f}" for i in jj)
    records.append(d)
df=pd.DataFrame(records).sort_values('original_motion_index')
df.to_csv(OUT/'analysis/verified_motion_analysis.csv',index=False)
stats={}
for k in ['cet_degrees',*rows[0]['GT_curl_descriptors']]:
    a=df[k].to_numpy(); stats[k]={'mean':float(a.mean()),'median':float(np.median(a)),'std_population':float(a.std()),'min':float(a.min()),'max':float(a.max()),'p10':float(np.percentile(a,10)),'p25':float(np.percentile(a,25)),'p75':float(np.percentile(a,75)),'p90':float(np.percentile(a,90)),'p95':float(np.percentile(a,95))}
stats['finger']={f:{'mean':float(df[f+'_cet'].mean()),'median':float(df[f+'_cet'].median()),'largest_count':int((df.top_finger==f).sum()),'ratio_of_means_pct':float(100*df[f+'_cet'].mean()/df[f+'_depth'].mean())} for f in summary['finger_names']}
stats['joints']=dict(zip(summary['joint_names'],np.mean([r['per_joint_degrees'] for r in rows],axis=0).tolist()))
stats['correlations']={k:{'pearson':float(df[k].corr(df.cet_degrees)),'spearman':float(df[k].rank().corr(df.cet_degrees.rank()))} for k in rows[0]['GT_curl_descriptors']}
stats['groups']=df.groupby('action').cet_degrees.agg(['count','mean','median']).sort_values('mean',ascending=False).reset_index().to_dict('records')
stats['unique_sequences']=df.sequence.nunique()
stats['hands']=df.hand_side.value_counts().to_dict()
stats['source_hashes']={f:hashlib.sha256((RUN/f).read_bytes()).hexdigest() for f in ['motions.json','motions.csv','summary.json']}
(OUT/'analysis/statistics.json').write_text(json.dumps(stats,indent=2))
doc=Document(ROOT/'CET_Curl_Error_with_Time_Alignment.docx')
(OUT/'analysis/source_document.txt').write_text('\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows),encoding='utf-8')
nb=json.loads((ROOT/'visualize_cet_results.ipynb').read_text(encoding='utf-8'))
(OUT/'analysis/notebook_source.txt').write_text('\n\n'.join(''.join(c['source']) for c in nb['cells']),encoding='utf-8')
for batch in range(3):
    files=sorted((ROOT/'exports/cet_graphs_20261006T025302388891Z').glob('*.png'))[batch*15:(batch+1)*15]
    sheet=Image.new('RGB',(1500,5*290),'white'); draw=ImageDraw.Draw(sheet)
    for i,f in enumerate(files):
        im=Image.open(f); im.thumbnail((490,260)); x=(i%3)*500; y=(i//3)*290
        sheet.paste(im,(x,y+25)); draw.text((x+5,y+4),f.name[:58],fill='black')
    sheet.save(OUT/f'analysis/export_contact_{batch+1}.jpg')
print(json.dumps(stats,indent=2))
cols=['original_motion_index','primitive','hand_side','cet_degrees','mean_absolute_flexion_degrees','flexion_range_degrees','total_angular_variation_degrees','variation_degrees_per_second','top_finger','little_relative_depth_pct']
for key in ['cet_degrees','mean_absolute_flexion_degrees','total_angular_variation_degrees','variation_degrees_per_second']:
    print('\nTOP',key); print(df.nlargest(12,key)[cols].to_string(index=False))
print('\nLOW ACTIVITY HIGH ERROR'); print(df.nsmallest(25,'variation_degrees_per_second').nlargest(6,'cet_degrees')[cols].to_string(index=False))
